#!/usr/bin/env python3
# © 2026 Jose Carlos Trejo Espinosa. Todos los derechos reservados.
"""Genera app/data/taxonomia.js, app/data/datos.js y app/img/ para la app del Túnel de Sanitizado 01
a partir del Excel y los .docx del repositorio.   Uso:  python3 tools/build_data.py   (requiere openpyxl y Pillow)"""
import json, re, zipfile, os, warnings, shutil
import openpyxl
from PIL import Image
warnings.filterwarnings('ignore')
AVISO = '/* © 2026 Jose Carlos Trejo Espinosa. Todos los derechos reservados. */\n'

REPO = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
APP = REPO + '/app'
XLSX = REPO + '/TUNEL SANITIZADO FORMATO.xlsx'
wb = openpyxl.load_workbook(XLSX, data_only=True)

def clean(v):
    if v is None: return ''
    return re.sub(r'\s+', ' ', str(v).replace('\xa0', ' ')).strip()

# ------------------------------------------------------------------ taxonomía
ws = wb['{TAXONOMIA.MAQUINA}']
rows = list(ws.iter_rows(values_only=True))
COMPS = [
    ('spray', 'SPREA (BOQUILLA DE ASPERSION)', 'mec', 'Boquilla de aspersión (Sprea)',
     'Boquilla hidráulica Unijet de ángulo ancho que rocía la solución sanitizante en niebla sobre el producto, dentro de la campana del túnel.'),
    ('motorreductor', 'MOTORREDUCTOR', 'mec', 'Motorreductor y electroválvula',
     'Motor trifásico WEG de 0.5 HP con reductor sin fin 40:1 (43.8 RPM) que mueve la banda, y solenoide de agua de 3/4" que abre el paso del agua.'),
    ('dossatron', 'BOMBA DOSSATRON', 'mec', 'Bomba dosificadora Dossatron',
     'Bomba dosificadora accionada por la presión del agua (D14MZ2-D). Succiona el sanitizante de la garrafa y lo mezcla con el agua al porcentaje ajustado.'),
    ('botonera', 'BOTON PARO DE EMERGENCIA (BOTONERA BM13)', 'mec', 'Botón de paro de emergencia (BM13)',
     'Botonera BM13 con el paro de emergencia, montada en una varilla con dos placas y tornillería inoxidable.'),
    ('transmision', 'TRANSMISION', 'mec', 'Transmisión, banda y estructura',
     'Estructura del túnel, banda modular Serie 900, sprockets, flechas, chumaceras, guías de desgaste, acoplamientos y tornillería.'),
    ('gabinete', 'GABINETE ELECTRICO', 'ele', 'Gabinete eléctrico',
     'Tablero de control con protecciones, contactor, guardamotor, variador de frecuencia, fuente, relevadores y clemas.'),
]
comp_by_name = {c[1]: c for c in COMPS}
taxo = {c[0]: {'id': c[0], 'sys': c[2], 'name': c[3], 'desc': c[4], 'items': []} for c in COMPS}
cur = None
for r in rows[16:92]:
    r = list(r) + [None] * 10
    if clean(r[1]) in comp_by_name:
        cur = comp_by_name[clean(r[1])][0]
    elif clean(r[0]) in comp_by_name:
        cur = comp_by_name[clean(r[0])][0]
    if cur is None: continue
    no, name = clean(r[2]), clean(r[3])
    if not name or not no or not re.match(r'^\d+$', no): continue
    q = clean(r[6])
    qty = int(q) if q.isdigit() else 1
    taxo[cur]['items'].append({
        'i': len(taxo[cur]['items']) + 1, 'no': int(no), 'n': name,
        'm': '' if clean(r[4]) in ('', 'N/A') else clean(r[4]),
        'p': '' if clean(r[5]) in ('', 'N/A') else clean(r[5]), 'q': qty})
# en planta las guías de desgaste (varillas blancas) son 4 y no 5 (lo confirmó el personal de planta)
for it in taxo['transmision']['items']:
    if it['i'] == 12: it['q_tax'] = it['q']; it['q'] = 4
# pendientes sin número de parte
for it in taxo['transmision']['items']:
    if it['p'].startswith('PENDIENTE'):
        it['pend'] = 1
        it['p'] = ('Sin No. de parte (obstruido en el escaneo)' if 'OBSTRUIDO' in it['p'] else 'Sin No. de parte visible en el manual')

# componentes que no están en la taxonomía del Excel pero salen en las rutinas, fotos y fallas
EXTRA = {
    'id': 'cubierta', 'sys': 'mec', 'name': 'Campana, cortinas, tina y caja del motor',
    'desc': 'Piezas que no figuran en la taxonomía del Excel pero aparecen en las tareas de las OT, las fotos de planta y el historial de fallas: campana del túnel, cortinas de PVC (hawaianas), tina (charola) de acero inoxidable, mangueras, garrafa de sanitizante y la caja de acero inoxidable que cubre el motorreductor en el lado B. Se modelaron a partir de las fotos de las rutinas y de lo que indicó el personal de planta.',
    'extra': 1,
    'items': [
        {'i': 1, 'no': 0, 'n': 'CAMPANA DEL TÚNEL (CUBIERTA DE ACERO INOXIDABLE)', 'm': 'ACERO INOXIDABLE', 'p': '', 'q': 1, 'x': 1},
        {'i': 2, 'no': 0, 'n': 'CORTINAS HAWAIANAS DE ENTRADA', 'm': 'PVC TRANSPARENTE', 'p': '', 'q': 7, 'x': 1},
        {'i': 3, 'no': 0, 'n': 'CORTINAS HAWAIANAS DE SALIDA', 'm': 'PVC TRANSPARENTE', 'p': '', 'q': 7, 'x': 1},
        {'i': 4, 'no': 0, 'n': 'TINA (CHAROLA) DE ACERO INOXIDABLE · TANQUE COLECTOR', 'm': 'ACERO INOXIDABLE', 'p': '', 'q': 1, 'x': 1},
        {'i': 5, 'no': 0, 'n': 'MANGUERAS DE DRENAJE', 'm': '', 'p': '', 'q': 2, 'x': 1},
        {'i': 6, 'no': 0, 'n': 'GARRAFA DE SANITIZANTE Y TUBO DE SUCCIÓN', 'm': 'POLIETILENO', 'p': '', 'q': 1, 'x': 1},
        {'i': 7, 'no': 0, 'n': 'MANGUERAS AZULES Y TRANSPARENTE, TUBERÍA DE LA CAMPANA Y CONDUIT', 'm': '', 'p': '', 'q': 1, 'x': 1},
        {'i': 8, 'no': 0, 'n': 'CAJA DE ACERO INOXIDABLE DEL MOTORREDUCTOR (LADO B)', 'm': 'ACERO INOXIDABLE', 'p': '', 'q': 1, 'x': 1},
    ]}
comps_out = [taxo[c[0]] for c in COMPS] + [EXTRA]
TAXO = {'sistemas': {'mec': 'Sistema mecánico', 'ele': 'Sistema eléctrico'}, 'comps': comps_out}
for c in comps_out:
    print(c['id'], len(c['items']))

# ------------------------------------------------------------------ tareas = las OT maestras (PDF)
# Las TAREAS del túnel son las de las dos órdenes de trabajo maestras (OT) en PDF: «Tarea No.», texto y horas estimadas.
# El Excel (hoja TAREAS) aporta lo que la OT no trae (sistema, táctica, requiere paro, nivel, herramientas) y las rutinas (.docx)
# aportan la foto de la ayuda visual; ambos se unen a cada tarea por el texto de la actividad (en las 22 tareas coincide al 100 %).
import subprocess, unicodedata
def norm(s):
    s = unicodedata.normalize('NFD', s or '').encode('ascii', 'ignore').decode().upper()
    return re.sub(r'[^A-Z0-9]+', ' ', s).strip()

def leer_ot(path):
    txt = subprocess.run(['pdftotext', '-layout', path, '-'], capture_output=True, text=True, check=True).stdout
    ot = {'orden': re.search(r'Orden No\.:\s*(\S+)', txt).group(1)}
    ot['descripcion'] = clean(re.search(r'Orden No\.:[^\n]*\n\s*(.+)', txt).group(1))
    ot['emitida'] = re.search(r'(\d\d)/(\d\d)/(\d{4})', txt).expand(r'\3-\1-\2')
    m = re.search(r'Equipo No\.[^\n]*\n\s*(OP \w+)\s+(.+?)\s{2,}(Qualtia.*)\n\s*(.*)', txt)
    ot['equipo_no'], ot['equipo'], ot['localizacion'] = m.group(1), clean(m.group(2)), clean(m.group(3) + ' ' + m.group(4))
    m = re.search(r'Prioridad[^\n]*\n\s*(\d+)\s+(\S+)\s+(\S+)', txt)
    ot['turno'], ot['tipo'], ot['clase'] = m.group(1), m.group(2), m.group(3)
    ot['mantenimiento'] = clean(re.search(r'ORDEN DE TRABAJO MAESTRA\s*\n\s*(.+)', txt).group(1))
    ot['tareas'], cur = [], None
    for ln in txt.split('\n'):
        m = re.match(r'\s*Tarea No:\s*(\S+)\s+Hrs Estim\.:\s*([\d.]+)', ln)
        if m: cur = {'codigo': m.group(1), 'hrs': float(m.group(2)), 'texto': ''}; ot['tareas'].append(cur); continue
        if cur is not None:
            if re.match(r'\s*Cada', ln): cur = None
            elif ln.strip(): cur['texto'] = (cur['texto'] + ' ' + clean(ln)).strip()
    return ot

OT_PDF = [('mensual', 'OT mensual', 'Inspección mensual general · equipo operando', REPO + '/OT TUNEL DE SANITIZADO 01 MENSUAL EQUIPO OPERANDO.pdf'),
          ('anual', 'OT anual', 'Mantenimiento preventivo mecánico-eléctrico · equipo parado', REPO + '/OT TUNEL DE SANITIZADO 01 ANUAL EQUIPO PARADO.pdf')]

# hoja TAREAS del Excel (por texto de la actividad)
EXCEL = {}
for r in wb['TAREAS '].iter_rows(min_row=8, values_only=True):
    per = clean(r[8])
    if not per.startswith('TS01'): continue
    EXCEL[norm(clean(r[5]))] = {'per_excel': per, 'sistema': clean(r[1]).capitalize(), 'tactica': clean(r[3]).upper(), 'frec': clean(r[6]).upper(),
                                'paro': clean(r[7]).upper(), 'nivel': clean(r[9]).upper(), 'min_excel': int(float(clean(r[10]) or 0)),
                                'herr': '' if clean(r[11]) in ('', 'N/A') else clean(r[11])}
print('tareas en el Excel:', len(EXCEL))

# pasos de las rutinas (.docx): número de paso de la ayuda visual, por texto
def docx_items(path):
    x = zipfile.ZipFile(path).read('word/document.xml').decode()
    paras = re.findall(r'<w:p[ >].*?</w:p>', x, re.S)
    out = []
    for p in paras:
        t = ''.join(re.findall(r'<w:t[^>]*>([^<]*)</w:t>', p)).strip()
        m = re.match(r'^(\d{4})\.-\s*\(\s*\)\s*(.+)$', t)
        if m: out.append((m.group(1), m.group(2).strip()))
    return out
PASO = {}      # texto normalizado -> número de paso de la rutina (1101…)
for p in ('RUTINA MENSUAL TUNEL DE SANITIZADO 01 EQUIPO OPERANDO.docx', 'RUTINA ANUAL TUNEL DE SANITIZADO 01 EQUIPO PARADO.docx'):
    for code, text in docx_items(REPO + '/' + p): PASO[norm(text)] = code
# sistema y clase salen del «Tarea No.» de la OT (TS01-<clase>-<sistema>-NN), que manda sobre las etiquetas del Excel
SISTEMA = {'MEC': 'Mecánico', 'ELEC': 'Eléctrico', 'GRAL': 'General'}
CLASE = {'INSP': 'Inspección', 'PRED': 'Predictivo', 'PREV': 'Preventivo', 'LIMP': 'Limpieza'}
# fotos de referencia (ver ref_XX.jpg) y puntos del 3D que cubre cada tarea (clave = paso de la rutina visual donde sale la foto)
FOTO = {'1101': 'ref_08', '1102': 'ref_11', '1103': 'ref_09', '1104': 'ref_01', '1105': 'ref_07', '1106': 'ref_08', '1107': 'ref_09',
        '1108': 'ref_11', '1109': 'ref_10', '1301': 'ref_09', '3101': 'ref_03', '3201': 'ref_09', '3302': 'ref_04', '7101': 'ref_06',
        '7102': 'ref_05', '1302': 'ref_a1', '1303': 'ref_a1', '1304': 'ref_a1', '1305': 'ref_a1', '1306': 'ref_a1', '1307': 'ref_a1', '3301': 'ref_04'}
PUNTOS = {
    '1101': ['transmision:12', 'transmision:15'], '1102': ['transmision:19'], '1103': ['transmision:35', 'transmision:34', 'transmision:22'],
    '1104': ['transmision:19'], '1105': ['dossatron:1', 'cubierta:6'], '1106': ['transmision:12', 'transmision:15'],
    '1107': ['transmision:20', 'transmision:6', 'transmision:32'], '1108': ['cubierta:2', 'cubierta:3'],
    '1109': ['spray:1', 'dossatron:1'], '1301': ['transmision:21', 'transmision:10', 'transmision:11'],
    '3101': ['comp:gabinete'], '3201': ['motorreductor:1', 'motorreductor:2'], '3302': ['comp:gabinete'],
    '7101': ['ALL'], '7102': ['transmision:1', 'transmision:12', 'transmision:15'],
    '1302': ['transmision:21'], '1303': ['motorreductor:2'], '1304': ['transmision:35', 'transmision:34'], '1305': ['motorreductor:2'],
    '1306': ['transmision:19'], '1307': ['motorreductor:1'], '3301': ['gabinete:5', 'gabinete:1', 'motorreductor:3']}
# notas que se muestran en la ruta guiada (respuestas del personal de planta)
NOTAS = {'1101': 'Según el personal de planta, los «rodillos y ruedas laterales» son las 4 varillas blancas de la entrada (foto 8).',
         '1302': 'En planta no hay una bomba de recirculación aparte (no sale en ninguna foto) y la foto de la ayuda visual anual señala la chumacera de la flecha; por eso se enlazan las chumaceras. Por confirmar.'}

import difflib
def buscar(dic, k):                     # coincidencia exacta o, si el texto cambia una letra, la más parecida (≥ 95 %)
    if k in dic: return dic[k]
    c = max(dic, key=lambda x: difflib.SequenceMatcher(None, x, k).ratio())
    return dic[c] if difflib.SequenceMatcher(None, c, k).ratio() >= 0.95 else None

ots = []
for oid, corto, titulo, path in OT_PDF:
    ot = leer_ot(path)
    items = []
    for t in ot['tareas']:
        k = norm(t['texto'])
        ex, paso = buscar(EXCEL, k), buscar(PASO, k)
        assert ex and paso, 'La tarea %s no se encontró por texto en el Excel o en la rutina' % t['codigo']
        items.append({'codigo': t['codigo'], 'rutina': paso, 'per_excel': ex['per_excel'], 'texto': t['texto'].lower(),
                      'sistema': SISTEMA[t['codigo'].split('-')[2]], 'clase': CLASE[t['codigo'].split('-')[1]], 'frec': ex['frec'], 'paro': ex['paro'], 'nivel': ex['nivel'],
                      'min': int(t['hrs']) if t['hrs'] == int(t['hrs']) else t['hrs'], 'herr': ex['herr'], 'foto': FOTO.get(paso, ''),
                      'puntos': PUNTOS[paso], 'nota': NOTAS.get(paso, '')})
        if ex['per_excel'] != t['codigo']: print('  aviso: la OT llama %s a la tarea que el Excel llama %s' % (t['codigo'], ex['per_excel']))
        if ex['min_excel'] != t['hrs']: print('  aviso: %s dura %s en la OT y %s en el Excel (se usa la OT)' % (t['codigo'], t['hrs'], ex['min_excel']))
    ots.append({'id': oid, 'corto': corto, 'titulo': titulo, 'orden': ot['orden'], 'descripcion': ot['descripcion'], 'equipo_no': ot['equipo_no'],
                'equipo': ot['equipo'], 'localizacion': ot['localizacion'], 'turno': ot['turno'], 'tipo': ot['tipo'], 'clase': ot['clase'],
                'mantenimiento': ot['mantenimiento'], 'emitida': ot['emitida'], 'items': items})
for o in ots: print(o['id'], o['orden'], len(o['items']), 'tareas')
assert len({it['codigo'] for o in ots for it in o['items']}) == 22

# ------------------------------------------------------------------ fallas
ws = wb['FALLAS REPETITIVAS']
rows = list(ws.iter_rows(values_only=True))
fallas, grp = [], None
for r in rows[5:]:
    r = list(r) + [None] * 40
    if clean(r[0]): grp = clean(r[0])
    if not clean(r[1]): continue
    def num(v):
        try: return float(clean(v) or 0)
        except ValueError: return 0.0
    fallas.append({'ot': clean(r[1]), 'desc': clean(r[3]), 'tipo': clean(r[9]), 'fecha': clean(r[10]), 'prio': clean(r[11]).split('.')[0],
                   'min': num(r[20]), 'muerto': num(r[21]), 'total': num(r[24]), 'clase': clean(r[26]), 'det': clean(r[29]), 'grupo': grp})
pareto = []
for r in rows[5:15]:
    r = list(r) + [None] * 40
    g = clean(r[33])
    if g and g != 'TOTAL': pareto.append({'grupo': g, 'ot': int(float(clean(r[34]))), 'pct': round(float(clean(r[35])) * 100, 1)})
print('fallas', len(fallas), 'pareto', len(pareto), sum(p['ot'] for p in pareto))
GRUPO_ELEM = {
    'BANDA - ROTURA / ESLABONES DAÑADOS': ['transmision:19'],
    'BANDA - ATORAMIENTO / DESALINEACION POR GUIAS': ['transmision:12', 'transmision:15', 'transmision:19'],
    'BANDA - FALLA DE TRANSMISION (SPROCKETS / CUÑA)': ['transmision:20', 'transmision:6', 'transmision:8', 'transmision:9'],
    'ESTRUCTURA - SOLDADURA DE GUIAS Y BASE (DESOLDADO)': ['transmision:1', 'transmision:15', 'transmision:12'],
    'SISTEMA ELECTRICO / GABINETE (NO ENCIENDE, PROTECCIONES, ELECTROVALVULA)': ['comp:gabinete', 'motorreductor:3'],
    'BANDA - CAMBIO DE HAWAIANAS (TOPES / ADITAMENTOS)': ['cubierta:2', 'cubierta:3'],
    'SPREA Y BOMBA DOSSATRON (SISTEMA DE ASPERSION)': ['spray:1', 'dossatron:1'],
}

# ------------------------------------------------------------------ refacciones
ws = wb['REFACCIONES ']
SEV = {'critico': 'CRÍTICO', 'moderado': 'MEDIA', 'normal': 'BAJA'}
REF_ELEM = [
    (r'SPREA', ['spray:1']), (r'SOLENOIDE', ['motorreductor:3']), (r'MOTOR ELECTRICO', ['motorreductor:1']), (r'REDUCTOR NEMA', ['motorreductor:2']),
    (r'DOSSATRON', ['dossatron:1']), (r'BOTON PARO', ['botonera:1']), (r'PLACA 1 BOTONERA', ['botonera:2']), (r'PLACA 2 BOTONERA', ['botonera:3']),
    (r'VARILLA BOTONERA', ['botonera:4']), (r'ESTRUCTURA TUNEL', ['transmision:1']), (r'CUÑA FLECHA', ['transmision:6']),
    (r'SOPORTE GUIA DE DESGASTE', ['transmision:15']), (r'GUIA DE DESGASTE', ['transmision:12']), (r'BANDA MODULAR', ['transmision:19']), (r'SPROCKET', ['transmision:20']),
    (r'HAWAIANAS', ['cubierta:2', 'cubierta:3']), (r'PERNO PLACA', ['transmision:2']), (r'PLACA SOPORTE ESTRUCTURA', ['transmision:3']),
    (r'FLECHA MOTOR', ['transmision:7']), (r'RODILLO CONDUCIDO', ['transmision:8']), (r'RODILLO MOTRIZ', ['transmision:9']),
    (r'RODILLO BANDA', ['transmision:17']), (r'SOPORTE RODILLO', ['transmision:18']), (r'CHUMACERA COCOL', ['transmision:21']),
    (r'COPLE MORDAZA \(MOTOR\)', ['transmision:22']), (r'COPLE MORDAZA \(TUNEL\)', ['transmision:34']), (r'ELEMENTO BUNA', ['transmision:35']),
    (r'PLACA 1 CHUMACERA', ['transmision:10']), (r'PLACA 2 CHUMACERA', ['transmision:11']),
    (r'SELENOIDE', ['transmision:33']), (r'GUARDA MOTOR', ['gabinete:5']), (r'VARIADOR', ['gabinete:10']), (r'TERMICO 3x10', ['gabinete:1']),
    (r'TERMICO 3x20', ['gabinete:2']), (r'CONTACTOR', ['gabinete:4']), (r'FUENTE', ['gabinete:6']), (r'RELEVADOR FINDER 220', ['gabinete:7']),
    (r'RELEVADOR FINDER 24', ['gabinete:8']), (r'RELEVADOR DE 14', ['gabinete:9']), (r'BOTON LATERAL', ['gabinete:3']),
    (r'TIERRA FISICA', ['gabinete:11']), (r'PASO 2\.5', ['gabinete:12']), (r'RIEL DIN', ['gabinete:13']), (r'CANALETA', ['gabinete:14']),
    (r'GLANDULA', ['gabinete:15'])]
refs = []
for r in ws.iter_rows(min_row=9, max_row=62, values_only=True):
    r = list(r) + [None] * 10
    sev = clean(r[7]).lower()
    if sev not in SEV: continue
    part = clean(r[0]); desc = clean(r[1])
    elems = []
    for pat, el in REF_ELEM:
        if re.search(pat, part + ' ' + desc, re.I): elems = el; break
    refs.append({'parte': part, 'desc': desc, 'uds': clean(r[2]), 'sev': SEV[sev], 'elems': elems})
print('refs', len(refs), 'sin elemento:', [x['parte'] for x in refs if not x['elems']])

# ------------------------------------------------------------------ escribir JS
os.makedirs(APP + '/data', exist_ok=True)
with open(APP + '/data/taxonomia.js', 'w', encoding='utf-8') as f:
    f.write(AVISO + 'window.TAXO=' + json.dumps(TAXO, ensure_ascii=False, separators=(',', ':')) + ';\n')
DATA = {'maquina': 'TÚNEL DE SANITIZADO 01 (OP TS01)', 'area': 'Empaque', 'linea': 'Cortes', 'ots': ots,
        'fallas': fallas, 'pareto': pareto, 'grupoElem': GRUPO_ELEM, 'refs': refs}
with open(APP + '/data/datos.js', 'w', encoding='utf-8') as f:
    f.write(AVISO + 'window.DATA=' + json.dumps(DATA, ensure_ascii=False, separators=(',', ':')) + ';\n')

# ------------------------------------------------------------------ imágenes
import io
IMG = APP + '/img'
os.makedirs(IMG, exist_ok=True)
def save(data, name, maxs=1000, q=84):
    im = Image.open(io.BytesIO(data))
    if im.mode in ('RGBA', 'LA', 'P'):
        im = im.convert('RGBA'); bg = Image.new('RGB', im.size, (255, 255, 255)); bg.paste(im, mask=im.split()[-1]); im = bg
    im = im.convert('RGB'); im.thumbnail((maxs, maxs)); im.save(IMG + '/' + name + '.jpg', quality=q, optimize=True)
zm = zipfile.ZipFile(REPO + '/REFERENCIA VISUAL MENSUAL (TUNEL DE SANITIZADO 01).docx')   # fotos de la rutina mensual
for i in range(1, 12):
    save(zm.read('word/media/image%d.jpeg' % i), 'ref_%02d' % i)
za = zipfile.ZipFile(REPO + '/REFERENCIA VISUAL ANUAL (TUNEL DE SANITIZADO 01).docx')    # foto de la rutina anual (chumacera y cople)
save(za.read('word/media/image1.jpeg'), 'ref_a1')
zx = zipfile.ZipFile(XLSX)                                                                # ilustraciones del manual (hoja de taxonomía)
for src, name in [('image49', 'x_spray'), ('image50', 'x_motorreductor'), ('image51', 'x_solenoide'), ('image52', 'x_dossatron'),
                  ('image53', 'x_botonera_desp'), ('image54', 'x_botonera_foto'), ('image55', 'x_transmision'),
                  ('image56', 'x_gabinete'), ('image57', 'x_general')]:
    save(zx.read('xl/media/' + src + '.png'), name, 1000, 88)
print(sorted(os.listdir(IMG)))
